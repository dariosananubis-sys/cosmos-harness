#!/usr/bin/env python3
"""gsheets.py — acceso a las hojas de Google que ve tu cuenta, sin re-login.

Usa el token guardado por `gcloud auth login --enable-gdrive-access` (persistente en
~/.config/gcloud). No pide acceso cada vez: mientras el login siga vivo, entra directo.

Uso:
  python3 gsheets.py list
  python3 gsheets.py read "<nombre o fileId>" [--tab NOMBRE] [--range A1:H50]
  python3 gsheets.py write <fileId> --tab NOMBRE --cell A1 --value "texto"   # solo Google Sheet nativa

Notas:
  - .xlsx (Excel en Drive) -> se descargan y leen con openpyxl (Sheets API no los abre).
  - write solo funciona en Google Sheets nativas y donde tu cuenta sea Editor.
"""
import argparse, io, os, subprocess, sys

try:
    import certifi
    os.environ.setdefault("SSL_CERT_FILE", certifi.where())
except Exception:
    pass

from google.oauth2.credentials import Credentials
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload

GSHEET = "application/vnd.google-apps.spreadsheet"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _token():
    try:
        return subprocess.check_output(["gcloud", "auth", "print-access-token"], text=True).strip()
    except Exception:
        sys.exit("Sin sesion gcloud. Ejecuta una vez:  gcloud auth login --enable-gdrive-access")


def _clients():
    creds = Credentials(token=_token())
    return (build("drive", "v3", credentials=creds, cache_discovery=False),
            build("sheets", "v4", credentials=creds, cache_discovery=False))


def _resolve(drive, ref):
    """Devuelve (fileId, name, mimeType) a partir de un id o un nombre."""
    got = None
    try:
        got = drive.files().get(fileId=ref, fields="id,name,mimeType").execute()
    except Exception:
        pass
    if got:
        return got["id"], got["name"], got["mimeType"]
    safe = ref.replace("'", "\\'")
    q = f"name contains '{safe}' and (mimeType='{GSHEET}' or mimeType='{XLSX}') and trashed=false"
    hits = drive.files().list(q=q, spaces="drive", fields="files(id,name,mimeType)", pageSize=10).execute().get("files", [])
    if not hits:
        sys.exit(f"No encuentro ninguna hoja que coincida con: {ref}")
    if len(hits) > 1:
        print("Varias coincidencias, usa el fileId exacto:", file=sys.stderr)
        for h in hits:
            print(f"  {h['id']}  {h['name']}", file=sys.stderr)
        sys.exit(1)
    h = hits[0]
    return h["id"], h["name"], h["mimeType"]


def cmd_list(drive, sheets, args):
    q = f"(mimeType='{GSHEET}' or mimeType='{XLSX}') and trashed=false"
    files, page = [], None
    while True:
        r = drive.files().list(q=q, spaces="drive",
            fields="nextPageToken, files(id,name,mimeType,owners(emailAddress),capabilities(canEdit))",
            pageSize=200, pageToken=page).execute()
        files.extend(r.get("files", [])); page = r.get("nextPageToken")
        if not page:
            break
    files.sort(key=lambda x: x["name"])
    print(f"{len(files)} hojas visibles:\n")
    for f in files:
        owner = f.get("owners", [{}])[0].get("emailAddress", "?")
        rw = "RW" if f.get("capabilities", {}).get("canEdit") else "R "
        kind = "GSheet" if f["mimeType"] == GSHEET else "xlsx  "
        print(f"  [{rw}] {kind}  {f['id']}  {f['name']}  (dueno: {owner})")


def _dump(rows):
    for row in rows:
        print("\t".join("" if c is None else str(c) for c in row))


def cmd_read(drive, sheets, args):
    fid, name, mime = _resolve(drive, args.ref)
    print(f"# {name}  ({fid})\n")
    if mime == GSHEET:
        meta = sheets.spreadsheets().get(spreadsheetId=fid, fields="sheets.properties.title").execute()
        tabs = [s["properties"]["title"] for s in meta.get("sheets", [])]
        tab = args.tab or tabs[0]
        rng = f"'{tab}'!{args.range}" if args.range else f"'{tab}'"
        vals = sheets.spreadsheets().values().get(spreadsheetId=fid, range=rng).execute().get("values", [])
        print(f"# pestanas: {tabs}\n# mostrando: {tab}\n")
        _dump(vals)
    else:
        import openpyxl
        buf = io.BytesIO(); dl = MediaIoBaseDownload(buf, drive.files().get_media(fileId=fid))
        done = False
        while not done:
            _, done = dl.next_chunk()
        buf.seek(0)
        wb = openpyxl.load_workbook(buf, read_only=True, data_only=True)
        tab = args.tab or wb.sheetnames[0]
        ws = wb[tab]
        print(f"# pestanas: {wb.sheetnames}\n# mostrando: {tab}\n")
        _dump(ws.iter_rows(values_only=True))


def cmd_write(drive, sheets, args):
    fid, name, mime = _resolve(drive, args.fileId)
    if mime != GSHEET:
        sys.exit("write solo soporta Google Sheets nativas (no .xlsx).")
    rng = f"'{args.tab}'!{args.cell}"
    sheets.spreadsheets().values().update(
        spreadsheetId=fid, range=rng, valueInputOption="USER_ENTERED",
        body={"values": [[args.value]]}).execute()
    print(f"OK: {name} -> {rng} = {args.value!r}")


def main():
    p = argparse.ArgumentParser(description="Acceso a Google Sheets/Drive de tu cuenta sin re-login")
    sub = p.add_subparsers(dest="cmd", required=True)
    sub.add_parser("list")
    r = sub.add_parser("read"); r.add_argument("ref"); r.add_argument("--tab"); r.add_argument("--range")
    w = sub.add_parser("write"); w.add_argument("fileId")
    w.add_argument("--tab", required=True); w.add_argument("--cell", required=True); w.add_argument("--value", required=True)
    args = p.parse_args()
    drive, sheets = _clients()
    {"list": cmd_list, "read": cmd_read, "write": cmd_write}[args.cmd](drive, sheets, args)


if __name__ == "__main__":
    main()
